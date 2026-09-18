"""
将已导出的课表 JSON 生成为包含整个学期各周课表的可视化 Excel 文件。

用法：
    python export_schedule_excel.py
    python export_schedule_excel.py selected_results/selected_courses_20260702_105748.json
    python export_schedule_excel.py --week-count 18 -o selected_results/course_schedule.xlsx

默认行为：
    - 自动读取 selected_results/ 下最新的 selected_courses_*.json；
    - 自动根据课表中的周次位串识别学期总周数；
    - 输出到 selected_results/course_schedule_all_weeks.xlsx；
    - 每个教学周单独占一个工作表，例如“第1周课表”“第2周课表”；
    - 最后一张表是全部课程明细，方便筛选和核对。

依赖：
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_RESULTS_DIR = BASE_DIR / "selected_results"
DAY_NAMES = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
DAY_NAME_TO_NUMBER = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "日": 7,
    "天": 7,
}
SECTION_GROUPS = (("上午", 1, 4), ("下午", 5, 8), ("晚上", 9, 13))


@dataclass(frozen=True)
class CourseSlot:
    name: str
    course_index: str
    week_name: str
    week_mask: str
    begin_section: int
    end_section: int
    day: int
    place: str
    teacher: str
    sport_name: str = ""

    @property
    def title(self) -> str:
        title = self.name.strip() or "未命名课程"
        if self.course_index.strip():
            title += f"-{self.course_index.strip()}"
        if self.sport_name.strip():
            title += f"({self.sport_name.strip()})"
        return title

    @property
    def section_text(self) -> str:
        return f"{self.week_name} {self.begin_section}-{self.end_section}节".strip()

    @property
    def display_text(self) -> str:
        return "\n".join(
            value
            for value in (self.title, self.section_text, self.place, self.teacher)
            if value
        )


@dataclass(frozen=True)
class DisplayBlock:
    day: int
    begin_section: int
    end_section: int
    text: str


def as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_int(value: Any) -> int | None:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def parse_day(value: Any) -> int | None:
    if value is None:
        return None
    text = as_text(value)
    number = parse_int(text)
    if number is not None and 1 <= number <= 7:
        return number
    for char, day in DAY_NAME_TO_NUMBER.items():
        if char in text:
            return day
    return None


def compact_week_ranges(mask: str) -> str:
    positions = [i + 1 for i, flag in enumerate(mask) if flag == "1"]
    if not positions:
        return ""

    ranges: list[str] = []
    start = previous = positions[0]
    for current in positions[1:]:
        if current == previous + 1:
            previous = current
            continue
        ranges.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = current
    ranges.append(str(start) if start == previous else f"{start}-{previous}")
    return ",".join(ranges) + "周"


def get_week_name(course: dict[str, Any], time_info: dict[str, Any]) -> str:
    week_name = as_text(time_info.get("weekName")) or as_text(course.get("weekName"))
    if week_name:
        return week_name

    week_mask = as_text(time_info.get("week")) or as_text(course.get("week"))
    return compact_week_ranges(week_mask)


def get_week_mask(course: dict[str, Any], time_info: dict[str, Any]) -> str:
    return as_text(time_info.get("week")) or as_text(course.get("week"))


def is_course_record(value: Any) -> bool:
    return isinstance(value, dict) and any(
        key in value for key in ("courseName", "teachingClassID", "beginSection")
    )


def iter_course_records(payload: Any) -> Iterable[dict[str, Any]]:
    """兼容 selected_courses JSON 和带 teachingTimeList 的课程查询 JSON。"""

    if isinstance(payload, list):
        for item in payload:
            yield from iter_course_records(item)
        return

    if not isinstance(payload, dict):
        return

    data_list = payload.get("dataList")
    if isinstance(data_list, list):
        for item in data_list:
            yield from iter_course_records(item)
        return

    if is_course_record(payload):
        time_list = payload.get("teachingTimeList")
        if isinstance(time_list, list) and time_list:
            for time_info in time_list:
                if isinstance(time_info, dict):
                    merged = dict(payload)
                    merged.update(time_info)
                    yield merged
            return
        yield payload
        return

    for value in payload.values():
        yield from iter_course_records(value)


def make_course_slot(record: dict[str, Any]) -> CourseSlot | None:
    day = parse_day(record.get("dayOfWeek"))
    begin = parse_int(record.get("beginSection"))
    end = parse_int(record.get("endSection"))
    if day is None or begin is None or end is None:
        return None
    if not 1 <= begin <= 13 or not 1 <= end <= 13:
        return None
    if end < begin:
        begin, end = end, begin

    return CourseSlot(
        name=as_text(record.get("courseName")) or "未命名课程",
        course_index=as_text(record.get("courseIndex")),
        week_name=get_week_name(record, record),
        week_mask=get_week_mask(record, record),
        begin_section=begin,
        end_section=end,
        day=day,
        place=as_text(record.get("teachingPlace")) or "时间地点未安排",
        teacher=as_text(record.get("teacherName")),
        sport_name=as_text(record.get("sportName")),
    )


def latest_input_file() -> Path:
    candidates = sorted(DEFAULT_RESULTS_DIR.glob("selected_courses_*.json"))
    if not candidates:
        raise FileNotFoundError(
            "未找到 selected_results/selected_courses_*.json，请先运行 query_selected.py，"
            "或在命令行中显式传入 JSON 路径。"
        )
    return candidates[-1]


def week_name_matches(week_name: str, week: int) -> bool:
    """当导出数据缺少 week 位串时，尽量从 weekName 判断周次。"""

    text = week_name.replace(" ", "")
    if not text:
        return True

    # API 正常会提供 week 位串；这里只作为兼容旧导出文件的兜底规则。
    if "单双周" in text:
        return True
    if "单周" in text and week % 2 == 0:
        return False
    if "双周" in text and week % 2 == 1:
        return False

    found_range = False
    for match in re.finditer(r"(\d+)\s*[-~至]\s*(\d+)", text):
        found_range = True
        if int(match.group(1)) <= week <= int(match.group(2)):
            return True
    if found_range:
        return False

    numbers = [int(number) for number in re.findall(r"\d+", text)]
    return not numbers or week in numbers


def is_active_in_week(slot: CourseSlot, week: int) -> bool:
    if slot.week_mask:
        return week <= len(slot.week_mask) and slot.week_mask[week - 1] == "1"
    return week_name_matches(slot.week_name, week)


def load_courses(input_path: Path) -> tuple[list[CourseSlot], str, str]:
    with input_path.open("r", encoding="utf-8") as file:
        payload = json.load(file)

    courses: list[CourseSlot] = []
    seen: set[tuple[Any, ...]] = set()
    school_term = ""
    student_code = ""
    for record in iter_course_records(payload):
        slot = make_course_slot(record)
        if slot is None:
            continue
        school_term = school_term or as_text(record.get("schoolTerm"))
        student_code = student_code or as_text(record.get("studentCode"))
        key = (
            slot.name,
            slot.course_index,
            slot.day,
            slot.begin_section,
            slot.end_section,
            slot.week_name,
            slot.place,
            slot.teacher,
        )
        if key not in seen:
            seen.add(key)
            courses.append(slot)

    return courses, school_term, student_code


def detect_week_count(courses: list[CourseSlot], default: int = 18) -> int:
    """从实际有课的最大周次推断学期长度，避免把位串尾部的 0 算成额外周次。"""

    max_week = 0
    for course in courses:
        if course.week_mask:
            active_weeks = [
                index + 1
                for index, flag in enumerate(course.week_mask)
                if flag == "1"
            ]
            if active_weeks:
                max_week = max(max_week, max(active_weeks))

        for match in re.finditer(r"(\d+)\s*[-~至]\s*(\d+)", course.week_name):
            max_week = max(max_week, int(match.group(2)))

    return max_week or default


def build_display_blocks(courses: list[CourseSlot]) -> list[DisplayBlock]:
    """合并同一天发生时间重叠的课程，避免 Excel 合并单元格相互冲突。"""

    blocks: list[DisplayBlock] = []
    for day in range(1, 8):
        day_courses = sorted(
            (course for course in courses if course.day == day),
            key=lambda course: (course.begin_section, course.end_section, course.title),
        )
        current: list[CourseSlot] = []
        current_end = 0
        for course in day_courses:
            if current and course.begin_section > current_end:
                blocks.append(
                    DisplayBlock(
                        day=day,
                        begin_section=min(item.begin_section for item in current),
                        end_section=current_end,
                        text="\n\n".join(item.display_text for item in current),
                    )
                )
                current = []
            current.append(course)
            current_end = max(current_end, course.end_section)

        if current:
            blocks.append(
                DisplayBlock(
                    day=day,
                    begin_section=min(item.begin_section for item in current),
                    end_section=current_end,
                    text="\n\n".join(item.display_text for item in current),
                )
            )
    return blocks


def fill(value: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=value)


def apply_border(ws: Any, min_row: int, min_col: int, max_row: int, max_col: int, border: Border) -> None:
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = border


def write_schedule_sheet(
    workbook: Workbook,
    courses: list[CourseSlot],
    week: int,
    school_term: str,
) -> None:
    sheet = workbook.create_sheet(f"第{week}周课表")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C5"

    dark_blue = "17345C"
    header_blue = "EAF2FC"
    group_blue = "EEF5FD"
    body_blue = "FBFDFD"
    border_blue = "D4E0F1"
    white = "FFFFFF"
    thin = Side(style="thin", color=border_blue)
    medium = Side(style="medium", color=dark_blue)
    light_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 列宽：A 为时段，B 为节次，C:I 为周一至周日。
    widths = {"A": 8, "B": 6, **{get_column_letter(column): 20 for column in range(3, 10)}}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A2:I2")
    sheet["A2"] = f"第{week}周课表"
    sheet["A2"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=dark_blue)
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 26

    sheet["A3"] = "学期"
    sheet["B3"] = school_term or "未知学期"
    sheet["D3"] = "显示周次"
    sheet["E3"] = f"第{week}周"
    sheet["G3"] = "课程数"
    sheet["H3"] = len(courses)
    for address in ("A3", "D3", "G3"):
        sheet[address].font = Font(name="Microsoft YaHei", size=10, bold=True, color=dark_blue)
    for address in ("B3", "E3", "H3"):
        sheet[address].font = Font(name="Microsoft YaHei", size=10, color=dark_blue)
    sheet["A3"].alignment = sheet["B3"].alignment = Alignment(vertical="center")
    sheet["D3"].alignment = sheet["E3"].alignment = Alignment(vertical="center")
    sheet["G3"].alignment = sheet["H3"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 20

    sheet.merge_cells("A4:B4")
    sheet["A4"] = "节次/星期"
    for column, day_name in enumerate(DAY_NAMES, start=3):
        sheet.cell(row=4, column=column, value=day_name)

    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color=dark_blue)
    header_alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[4][0:9]:
        cell.fill = fill(header_blue)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = light_border
    sheet.row_dimensions[4].height = 25

    # 主体 13 节。
    for section in range(1, 14):
        row = 4 + section
        sheet.cell(row=row, column=2, value=section)
        for column in range(1, 10):
            cell = sheet.cell(row=row, column=column)
            cell.fill = fill(body_blue)
            cell.font = Font(name="Microsoft YaHei", size=10, color=dark_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = light_border
        sheet.row_dimensions[row].height = 43

    for group_name, start, end in SECTION_GROUPS:
        start_row = 4 + start
        end_row = 4 + end
        sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        cell = sheet.cell(row=start_row, column=1)
        cell.value = group_name
        cell.fill = fill(group_blue)
        cell.font = Font(name="Microsoft YaHei", size=10, color=dark_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        apply_border(sheet, start_row, 1, end_row, 1, light_border)

    # 将指定周的课程写入对应日期和节次的合并单元格。
    for block in build_display_blocks(courses):
        start_row = 4 + block.begin_section
        end_row = 4 + block.end_section
        column = 2 + block.day
        sheet.merge_cells(start_row=start_row, start_column=column, end_row=end_row, end_column=column)
        cell = sheet.cell(row=start_row, column=column)
        cell.value = block.text
        cell.fill = fill(white)
        cell.font = Font(name="Microsoft YaHei", size=9, color=dark_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        apply_border(sheet, start_row, column, end_row, column, light_border)

    sheet.merge_cells("A19:I19")
    note = "说明：本页仅显示该周实际有课的课程；课程块内依次为课程、周次/节次、地点、教师。"
    if not courses:
        note = "说明：本周没有检测到课程，请确认原始课表 JSON。"
    sheet["A19"] = note
    sheet["A19"].font = Font(name="Microsoft YaHei", size=9, italic=True, color="5A6B82")
    sheet["A19"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[19].height = 22

    sheet.print_area = "A1:I19"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    sheet.page_margins = PageMargins(left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.1, footer=0.1)
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.tabColor = dark_blue
    apply_border(sheet, 4, 1, 17, 9, light_border)
    sheet["A4"].border = Border(left=medium, top=medium, bottom=thin)
    sheet["I17"].border = Border(right=medium, bottom=medium)


def write_detail_sheet(
    workbook: Workbook,
    courses: list[CourseSlot],
) -> None:
    sheet = workbook.create_sheet("全部课程明细")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    headers = ["星期", "开始节次", "结束节次", "课程", "上课周次", "地点", "教师"]
    rows = [headers]
    for course in sorted(courses, key=lambda item: (item.day, item.begin_section, item.end_section, item.title)):
        rows.append(
            [
                DAY_NAMES[course.day - 1],
                course.begin_section,
                course.end_section,
                course.title,
                course.week_name,
                course.place,
                course.teacher,
            ]
        )

    for row_index, row_values in enumerate(rows, start=1):
        for column_index, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = Font(name="Microsoft YaHei", size=10, color="17345C", bold=row_index == 1)
            cell.alignment = Alignment(horizontal="center" if row_index == 1 else "left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="D4E0F1"))
            if row_index == 1:
                cell.fill = fill("EAF2FC")
        sheet.row_dimensions[row_index].height = 24 if row_index == 1 else 34

    widths = {"A": 10, "B": 12, "C": 12, "D": 34, "E": 16, "F": 22, "G": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:G{max(1, len(rows))}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
    sheet.sheet_properties.tabColor = "7F8FA6"


def export_workbook(
    output_path: Path,
    courses: list[CourseSlot],
    week_count: int,
    school_term: str,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "学期课程表"
    workbook.properties.subject = "暨南大学教务系统课表"
    workbook.properties.creator = "JNU Course Schedule Exporter"
    for week in range(1, week_count + 1):
        week_courses = [course for course in courses if is_active_in_week(course, week)]
        write_schedule_sheet(workbook, week_courses, week, school_term)
    write_detail_sheet(workbook, courses)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="生成包含整个学期各周课表的可视化 Excel 文件")
    parser.add_argument(
        "input_json",
        nargs="?",
        type=Path,
        help="课表 JSON；省略时读取 selected_results/ 下最新的 selected_courses_*.json",
    )
    parser.add_argument(
        "--week-count",
        type=int,
        help="学期总周数；省略时根据课表中实际有课的最大周次自动识别",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="输出 XLSX 路径；省略时输出到 selected_results/course_schedule_all_weeks.xlsx",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        input_path = (args.input_json or latest_input_file()).resolve()
        if not input_path.exists():
            raise FileNotFoundError(f"找不到输入 JSON：{input_path}")

        output_path = args.output or (DEFAULT_RESULTS_DIR / "course_schedule_all_weeks.xlsx")
        output_path = output_path.resolve()
        courses, school_term, _student_code = load_courses(input_path)
        week_count = args.week_count or detect_week_count(courses)
        if not 1 <= week_count <= 52:
            raise ValueError("学期总周数必须在 1 到 52 之间。")
        export_workbook(output_path, courses, week_count, school_term)

        print(f"已导出 {week_count} 周课表：{output_path}")
        print(f"原始课程数：{len(courses)}")
        return 0
    except (FileNotFoundError, json.JSONDecodeError, OSError, ValueError) as error:
        print(f"导出失败：{error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
